"""
tests/backtesting/test_statistical_tests.py
==========================================

Unit Test Suite for Statistical Significance & Hypothesis Testing Engine (Part 9.5).
"""

import json
from pathlib import Path
import openpyxl
import numpy as np
import pandas as pd
import pytest

from backtesting.statistical_tests import (
    BootstrapCIResult,
    DieboldMarianoResult,
    EffectSizeResult,
    PairedTestResult,
    StatisticalTestEngine,
)


@pytest.fixture
def synthetic_statistical_data() -> tuple[np.ndarray, dict[str, np.ndarray], dict[str, np.ndarray]]:
    np.random.seed(42)
    n_hours = 240  # 10 days
    ground_truth = 30.0 + 10.0 * np.sin(np.linspace(0, 10 * 2 * np.pi, n_hours))

    # Model A is accurate (noise = 1.5), Model B is noisy (noise = 6.0)
    pred_a = ground_truth + np.random.normal(0, 1.5, n_hours)
    pred_b = ground_truth + np.random.normal(0, 6.0, n_hours)

    models_pred = {
        "Accurate_ML": pred_a,
        "Noisy_Baseline": pred_b,
    }

    daily_rev = {
        "Accurate_ML": np.array([2500.0, 2600.0, 2400.0, 2550.0, 2700.0, 2450.0, 2650.0, 2500.0, 2550.0, 2600.0]),
        "Noisy_Baseline": np.array([1800.0, 1900.0, 1750.0, 1850.0, 2000.0, 1700.0, 1950.0, 1800.0, 1850.0, 1900.0]),
    }

    return ground_truth, models_pred, daily_rev


class TestStatisticalTestsMath:
    def test_diebold_mariano_superiority_detection(self, synthetic_statistical_data):
        ground_truth, models_pred, _ = synthetic_statistical_data
        dm_res = StatisticalTestEngine.diebold_mariano_test(
            y_true=ground_truth,
            y_pred_a=models_pred["Accurate_ML"],
            y_pred_b=models_pred["Noisy_Baseline"],
            forecast_horizon_h=24,
            model_a_name="Accurate_ML",
            model_b_name="Noisy_Baseline",
        )

        assert dm_res.mean_loss_a < dm_res.mean_loss_b
        assert dm_res.hln_statistic < 0.0  # Loss A - Loss B < 0
        assert dm_res.p_value < 0.05
        assert dm_res.is_statistically_significant
        assert dm_res.superior_model == "Accurate_ML"

    def test_identical_models_dm_test(self, synthetic_statistical_data):
        ground_truth, models_pred, _ = synthetic_statistical_data
        dm_res = StatisticalTestEngine.diebold_mariano_test(
            y_true=ground_truth,
            y_pred_a=models_pred["Accurate_ML"],
            y_pred_b=models_pred["Accurate_ML"],
            model_a_name="Model_1",
            model_b_name="Model_2",
        )

        assert dm_res.mean_differential == 0.0
        assert dm_res.hln_statistic == 0.0
        assert dm_res.p_value == 1.0
        assert not dm_res.is_statistically_significant
        assert dm_res.superior_model == "Indistinguishable"

    def test_paired_tests_significance(self, synthetic_statistical_data):
        _, _, daily_rev = synthetic_statistical_data
        res = StatisticalTestEngine.paired_difference_tests(
            daily_rev["Accurate_ML"],
            daily_rev["Noisy_Baseline"],
            comparison_name="ML_vs_Baseline",
        )

        assert res.mean_difference > 0.0
        assert res.t_statistic > 0.0
        assert res.t_p_value < 0.01
        assert res.is_t_significant
        assert res.is_wilcoxon_significant

    def test_bootstrap_confidence_interval_coverage(self):
        np.random.seed(42)
        true_mean = 50.0
        sample = np.random.normal(true_mean, 5.0, 100)

        res = StatisticalTestEngine.bootstrap_confidence_interval(
            data=sample,
            statistic_fn=np.mean,
            n_bootstraps=1000,
            metric_name="Sample_Mean",
        )

        assert res.ci_95_lower <= res.point_estimate <= res.ci_95_upper
        assert res.ci_99_lower <= res.ci_95_lower
        assert res.ci_99_upper >= res.ci_95_upper
        assert abs(res.point_estimate - true_mean) < 2.0

    def test_effect_size_magnitude(self, synthetic_statistical_data):
        _, _, daily_rev = synthetic_statistical_data
        res = StatisticalTestEngine.calculate_effect_sizes(
            daily_rev["Accurate_ML"],
            daily_rev["Noisy_Baseline"],
            comparison_name="ML_vs_Baseline",
        )

        # 2550 vs 1850 is a large shift
        assert res.hedges_g > 1.5
        assert res.magnitude_interpretation == "Large"
        assert res.cles_pct == 100.0  # Every day Accurate_ML beat Noisy_Baseline


class TestStatisticalTestsPipeline:
    def test_evaluate_and_export_artifacts(self, synthetic_statistical_data, tmp_path):
        ground_truth, models_pred, daily_rev = synthetic_statistical_data
        engine = StatisticalTestEngine(output_directory=tmp_path)
        dm_res, boot_res, effect_res, artifacts = engine.evaluate_and_export(
            ground_truth=ground_truth,
            models_predictions=models_pred,
            daily_revenues=daily_rev,
        )

        assert artifacts.summary_csv.exists()
        assert artifacts.bootstrap_csv.exists()
        assert artifacts.dm_matrix_csv.exists()
        assert artifacts.statistical_excel.exists()
        assert artifacts.summary_json.exists()

        # Excel Verification
        wb = openpyxl.load_workbook(artifacts.statistical_excel, read_only=True)
        assert "Diebold_Mariano" in wb.sheetnames
        assert "Bootstrap_CIs" in wb.sheetnames
        assert "Effect_Sizes" in wb.sheetnames
        wb.close()

        # Figures Verification
        figures = list(artifacts.figures_directory.glob("*.png"))
        assert len(figures) == 5
        figure_names = {f.name for f in figures}
        assert figure_names == {
            "bootstrap_confidence_intervals.png",
            "diebold_mariano_matrix.png",
            "loss_differential_timeseries.png",
            "paired_difference_distribution.png",
            "effect_size_summary.png",
        }