"""
tests/backtesting/test_sensitivity_analysis.py
=============================================

Unit Test Suite for Sensitivity Analysis & Parametric Elasticity Engine (Part 9.4).
"""

import json
from pathlib import Path
import openpyxl
import numpy as np
import pandas as pd
import pytest

from backtesting.sensitivity_analysis import (
    ElasticityRecord,
    SensitivityAnalysisEngine,
    TornadoParameterRecord,
)


@pytest.fixture
def sample_scenarios_dataset() -> pd.DataFrame:
    return pd.DataFrame([
        # Horizon variations
        {"scenario_name": "horizon_12h", "category": "Forecast_Horizon", "forecast_horizon": 12, "net_revenue_usd": 700000.0, "gross_revenue_usd": 850000.0, "degradation_cost_usd": 150000.0, "final_soh": 0.985, "cumulative_efc": 160.0, "sharpe_ratio": 2.0},
        {"scenario_name": "horizon_24h", "category": "Forecast_Horizon", "forecast_horizon": 24, "net_revenue_usd": 780000.0, "gross_revenue_usd": 950000.0, "degradation_cost_usd": 170000.0, "final_soh": 0.983, "cumulative_efc": 175.0, "sharpe_ratio": 2.2},
        {"scenario_name": "horizon_48h", "category": "Forecast_Horizon", "forecast_horizon": 48, "net_revenue_usd": 848333.0, "gross_revenue_usd": 1102000.0, "degradation_cost_usd": 253667.0, "final_soh": 0.981, "cumulative_efc": 190.0, "sharpe_ratio": 2.4},
        {"scenario_name": "horizon_72h", "category": "Forecast_Horizon", "forecast_horizon": 72, "net_revenue_usd": 890000.0, "gross_revenue_usd": 1180000.0, "degradation_cost_usd": 290000.0, "final_soh": 0.979, "cumulative_efc": 200.0, "sharpe_ratio": 2.6},
        # Thermal variations
        {"scenario_name": "temp_15c_subcooled", "category": "Thermal_Sensitivity", "forecast_horizon": 48, "net_revenue_usd": 880000.0, "gross_revenue_usd": 1102000.0, "degradation_cost_usd": 222000.0, "final_soh": 0.987, "cumulative_efc": 190.0, "sharpe_ratio": 2.5},
        {"scenario_name": "temp_25c_reference", "category": "Thermal_Sensitivity", "forecast_horizon": 48, "net_revenue_usd": 848333.0, "gross_revenue_usd": 1102000.0, "degradation_cost_usd": 253667.0, "final_soh": 0.981, "cumulative_efc": 190.0, "sharpe_ratio": 2.4},
        {"scenario_name": "temp_45c_severe_stress", "category": "Thermal_Sensitivity", "forecast_horizon": 48, "net_revenue_usd": 680000.0, "gross_revenue_usd": 1102000.0, "degradation_cost_usd": 422000.0, "final_soh": 0.965, "cumulative_efc": 190.0, "sharpe_ratio": 1.9},
        # Efficiency variations
        {"scenario_name": "eff_85pct_aged", "category": "Efficiency_Sensitivity", "forecast_horizon": 48, "net_revenue_usd": 740000.0, "gross_revenue_usd": 980000.0, "degradation_cost_usd": 240000.0, "final_soh": 0.981, "cumulative_efc": 170.0, "sharpe_ratio": 2.1},
        {"scenario_name": "eff_90pct_baseline", "category": "Efficiency_Sensitivity", "forecast_horizon": 48, "net_revenue_usd": 848333.0, "gross_revenue_usd": 1102000.0, "degradation_cost_usd": 253667.0, "final_soh": 0.981, "cumulative_efc": 190.0, "sharpe_ratio": 2.4},
        {"scenario_name": "eff_95pct_nextgen", "category": "Efficiency_Sensitivity", "forecast_horizon": 48, "net_revenue_usd": 930000.0, "gross_revenue_usd": 1210000.0, "degradation_cost_usd": 280000.0, "final_soh": 0.981, "cumulative_efc": 205.0, "sharpe_ratio": 2.7},
        # Chemistry variations
        {"scenario_name": "chem_nmc_baseline", "category": "Battery_Chemistry", "forecast_horizon": 48, "net_revenue_usd": 848333.0, "gross_revenue_usd": 1102000.0, "degradation_cost_usd": 253667.0, "final_soh": 0.981, "cumulative_efc": 190.0, "sharpe_ratio": 2.4},
        {"scenario_name": "chem_lfp_stationary", "category": "Battery_Chemistry", "forecast_horizon": 48, "net_revenue_usd": 932814.0, "gross_revenue_usd": 1080000.0, "degradation_cost_usd": 147186.0, "final_soh": 0.989, "cumulative_efc": 185.0, "sharpe_ratio": 2.7},
        {"scenario_name": "chem_lto_heavy_cycle", "category": "Battery_Chemistry", "forecast_horizon": 48, "net_revenue_usd": 994495.0, "gross_revenue_usd": 1060000.0, "degradation_cost_usd": 65505.0, "final_soh": 0.995, "cumulative_efc": 180.0, "sharpe_ratio": 2.9},
        # Sizing variations
        {"scenario_name": "size_25mw_50mwh", "category": "System_Sizing", "forecast_horizon": 48, "net_revenue_usd": 424166.0, "gross_revenue_usd": 551000.0, "degradation_cost_usd": 126834.0, "final_soh": 0.981, "cumulative_efc": 190.0, "sharpe_ratio": 2.4},
        {"scenario_name": "size_100mw_200mwh", "category": "System_Sizing", "forecast_horizon": 48, "net_revenue_usd": 1696600.0, "gross_revenue_usd": 2204000.0, "degradation_cost_usd": 507400.0, "final_soh": 0.981, "cumulative_efc": 190.0, "sharpe_ratio": 2.4},
    ])


class TestSensitivityAnalysisMath:
    def test_arc_elasticity_formula(self):
        e = SensitivityAnalysisEngine.calculate_arc_elasticity(
            x_base=100.0, x_var=110.0, y_base=50.0, y_var=60.0
        )
        assert e > 0.0
        assert np.isclose(e, 1.909, atol=0.01)

    def test_zero_delta_elasticity_safety(self):
        e = SensitivityAnalysisEngine.calculate_arc_elasticity(
            x_base=100.0, x_var=100.0, y_base=50.0, y_var=60.0
        )
        assert e == 0.0

    def test_tornado_ranking_order(self, sample_scenarios_dataset):
        engine = SensitivityAnalysisEngine()
        tornado_records = engine.evaluate_tornado(sample_scenarios_dataset)

        assert len(tornado_records) >= 4
        assert tornado_records[0].parameter == "System Duration / Sizing"
        assert tornado_records[0].sensitivity_rank == 1

        for i in range(len(tornado_records) - 1):
            assert tornado_records[i].swing_usd >= tornado_records[i + 1].swing_usd

    def test_elasticity_evaluation_contents(self, sample_scenarios_dataset):
        engine = SensitivityAnalysisEngine()
        elasticities = engine.evaluate_elasticities(sample_scenarios_dataset)

        dimensions = {e.parameter_dimension for e in elasticities}
        assert "Forecast_Horizon" in dimensions
        assert "Thermal_Sensitivity" in dimensions
        assert "Efficiency_Sensitivity" in dimensions

        for e in elasticities:
            assert np.isfinite(e.arc_elasticity)
            assert np.isfinite(e.marginal_rate_usd)


class TestSensitivityAnalysisPipeline:
    def test_evaluate_and_export_artifacts(self, sample_scenarios_dataset, tmp_path):
        engine = SensitivityAnalysisEngine(output_directory=tmp_path)
        elasticities, tornado_records, artifacts = engine.evaluate_and_export(sample_scenarios_dataset)

        assert artifacts.summary_csv.exists()
        assert artifacts.elasticity_csv.exists()
        assert artifacts.tornado_csv.exists()
        assert artifacts.sensitivity_excel.exists()
        assert artifacts.summary_json.exists()

        wb = openpyxl.load_workbook(artifacts.sensitivity_excel, read_only=True)
        assert "Tornado_Spectrum" in wb.sheetnames
        assert "Elasticities" in wb.sheetnames
        wb.close()

        with open(artifacts.summary_json, "r", encoding="utf-8") as f:
            data = json.load(f)
        assert "most_sensitive_parameter" in data
        assert "tornado_spectrum" in data

        figures = list(artifacts.figures_directory.glob("*.png"))
        assert len(figures) == 5
        figure_names = {f.name for f in figures}
        assert figure_names == {
            "tornado_sensitivity_chart.png",
            "spider_radar_chart.png",
            "horizon_elasticity_curve.png",
            "thermal_degradation_surface.png",
            "efficiency_revenue_elasticity.png",
        }